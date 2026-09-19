import hashlib
import io
import re
import secrets
import unicodedata
import zipfile

from django.core.cache import cache
from django.db import transaction
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from contacts.models import Contact, ContactCategory
from lead_collector.models import WhatsAppContactImportRecord
from lead_collector.services.identity import normalized_phone
from lead_collector.services.whatsapp_contact_import import ensure_contact_categories
from organizations.models import Organization

MAX_FILE_SIZE = 8 * 1024 * 1024
MAX_SHEETS = 25
MAX_ROWS = 10000
MAX_COLUMNS = 100
MAX_UNCOMPRESSED_SIZE = 64 * 1024 * 1024
MAX_CELL_LENGTH = 10000
SESSION_TTL = 30 * 60
FIELDS = ("whatsapp_number", "business_name", "category", "email", "website", "address", "whatsapp_status", "evidence_type", "evidence_url", "provider")
ALIASES = {
    "whatsapp_number": ("normalized whatsapp number", "whatsapp number", "whatsapp", "mobile", "phone", "contact number"),
    "business_name": ("business name", "company", "company name", "name", "hotel name"),
    "category": ("category", "type"), "email": ("email", "email address"),
    "website": ("website", "website url"), "address": ("address",),
    "whatsapp_status": ("whatsapp status", "status"), "evidence_type": ("evidence type",),
    "evidence_url": ("evidence url", "source url"), "provider": ("provider", "source"),
}


class WorkbookError(ValueError): pass


def _text(value):
    if value is None: return ""
    if isinstance(value, str) and value.startswith("="): return ""
    return str(value).strip()[:MAX_CELL_LENGTH]


def _norm(value):
    return re.sub(r"[^a-z0-9]+", " ", unicodedata.normalize("NFKC", _text(value)).casefold()).strip()


def _open(data):
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            entries = archive.infolist()
            if len(entries) > 5000 or sum(item.file_size for item in entries) > MAX_UNCOMPRESSED_SIZE:
                raise WorkbookError("The workbook expands beyond the safe processing limit.")
            if any(item.flag_bits & 0x1 for item in entries):
                raise WorkbookError("Password-protected workbooks are not supported.")
    except zipfile.BadZipFile as exc:
        raise WorkbookError("The workbook is corrupt or cannot be read safely.") from exc
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True, keep_links=False)
    except (InvalidFileException, OSError, ValueError, KeyError, EOFError) as exc:
        raise WorkbookError("The workbook is corrupt or cannot be read safely.") from exc
    if not workbook.sheetnames: raise WorkbookError("The workbook contains no worksheets.")
    if len(workbook.sheetnames) > MAX_SHEETS: raise WorkbookError(f"Workbooks may contain at most {MAX_SHEETS} worksheets.")
    return workbook


def _sheet(workbook, name):
    if name not in workbook.sheetnames: raise WorkbookError("The selected worksheet was not found.")
    ws = workbook[name]
    if ws.max_row > MAX_ROWS or ws.max_column > MAX_COLUMNS: raise WorkbookError(f"A worksheet may contain at most {MAX_ROWS} rows and {MAX_COLUMNS} columns.")
    return ws


def _extract(ws):
    values = [[_text(cell) for cell in row[:MAX_COLUMNS]] for row in ws.iter_rows(values_only=True, max_row=MAX_ROWS + 1)]
    useful = [(index + 1, row) for index, row in enumerate(values) if any(row)]
    if not useful: raise WorkbookError("The selected worksheet is empty.")
    best = max(useful[:10], key=lambda item: sum(bool(v) for v in item[1]))
    header_row, raw_headers = best
    headers, seen = [], {}
    for index, value in enumerate(raw_headers):
        name = value or f"Column {index + 1}"
        seen[name] = seen.get(name, 0) + 1
        headers.append(name if seen[name] == 1 else f"{name} ({seen[name]})")
    rows = [{headers[i]: value for i, value in enumerate(row[:len(headers)])} | {"_row": number} for number, row in useful if number > header_row]
    return header_row, headers, rows


def suggestions(headers):
    result = {}
    for field, aliases in ALIASES.items():
        matches = [header for header in headers if _norm(header) in aliases]
        if matches: result[field] = matches[0]
    return result


def store_workbook(user, organization, uploaded):
    if not uploaded.name.lower().endswith(".xlsx"): raise WorkbookError("Only .xlsx workbooks are supported.")
    if uploaded.size > MAX_FILE_SIZE: raise WorkbookError("The workbook exceeds the 8 MB upload limit.")
    data = uploaded.read(MAX_FILE_SIZE + 1)
    if len(data) > MAX_FILE_SIZE: raise WorkbookError("The workbook exceeds the 8 MB upload limit.")
    workbook = _open(data)
    names = workbook.sheetnames
    recognized = next((name for name in names if _norm(name) == "whatsapp contacts"), None)
    selected = recognized or (names[0] if len(names) == 1 else None)
    token = secrets.token_urlsafe(32)
    cache.set(f"excel-wa:{token}", {"user": user.pk, "organization": organization.pk, "filename": uploaded.name.rsplit("/", 1)[-1].rsplit("\\", 1)[-1], "data": data}, SESSION_TTL)
    result = {"token": token, "sheets": names, "selected_sheet": selected, "expires_in": SESSION_TTL}
    if selected: result["sheet"] = inspect_sheet(token, user, organization, selected)
    return result


def session(token, user, organization):
    value = cache.get(f"excel-wa:{token}")
    if not value or value["user"] != user.pk or value["organization"] != organization.pk: raise WorkbookError("This upload session is invalid or has expired. Please upload the workbook again.")
    return value


def inspect_sheet(token, user, organization, sheet_name):
    value = session(token, user, organization); workbook = _open(value["data"]); ws = _sheet(workbook, sheet_name)
    header_row, headers, rows = _extract(ws)
    return {"name": sheet_name, "header_row": header_row, "headers": headers, "rows": rows[:15], "suggested_mappings": suggestions(headers)}


def _category_map(organization):
    ensure_contact_categories(organization)
    return {_norm(item.name): item for item in ContactCategory.objects.filter(organization=organization)}


def processed_rows(token, user, organization, sheet_name, mappings, default_category_id, confirmed):
    if not confirmed: raise WorkbookError("Confirm that the mapped numbers are authorized WhatsApp contacts.")
    if not mappings.get("whatsapp_number"): raise WorkbookError("WhatsApp Number must be mapped.")
    used = [v for v in mappings.values() if v]
    if len(used) != len(set(used)): raise WorkbookError("An Excel column cannot be mapped to more than one CRM field.")
    value = session(token, user, organization); workbook = _open(value["data"]); _, headers, raw = _extract(_sheet(workbook, sheet_name))
    if any(column not in headers for column in used): raise WorkbookError("One or more mapped columns are invalid.")
    categories = _category_map(organization)
    try: default = ContactCategory.objects.get(pk=default_category_id, organization=organization)
    except ContactCategory.DoesNotExist as exc: raise WorkbookError("Select a valid default category.") from exc
    seen, rows = {}, []
    for source in raw:
        number_value = source.get(mappings["whatsapp_number"], "")
        number = normalized_phone(number_value)
        category_text = source.get(mappings.get("category"), "") if mappings.get("category") else ""
        category = categories.get(_norm(category_text), default)
        item = {field: source.get(mappings.get(field), "") if mappings.get(field) else "" for field in FIELDS}
        item.update({"row": source["_row"], "number": number, "category_id": category.id, "category": category.name, "category_fallback": bool(category_text and _norm(category_text) not in categories)})
        if not number_value: item["status"] = "MISSING_NUMBER"
        elif not number or len(re.sub(r"\D", "", number)) < 7: item["status"] = "INVALID_NUMBER"
        elif number in seen:
            item["status"] = "DUPLICATE_IN_FILE"
            first = seen[number]
            for field in ("business_name", "email", "website", "address", "evidence_type", "evidence_url", "provider"):
                if not first[field] and item[field]: first[field] = item[field]
            if first["category"] == "Unknown" and item["category"] != "Unknown": first.update(category=item["category"], category_id=item["category_id"])
        else:
            existing = next((c for c in Contact.objects.filter(organization=organization).select_related("category") if normalized_phone(c.phone_number) == number), None)
            item["status"] = "READY" if not existing else "CATEGORY_UPDATE" if (not existing.category or existing.category.name == "Unknown") and category.name != "Unknown" else "ALREADY_EXISTS"
            item["contact_id"] = existing.id if existing else None; seen[number] = item
        rows.append(item)
    return value, rows


def preview(*args, **kwargs):
    _, rows = processed_rows(*args, **kwargs)
    counts = {"rows": len(rows), "valid": 0, "new": 0, "existing": 0, "duplicates": 0, "invalid": 0, "category_updates": 0, "skipped": 0}
    for row in rows:
        status = row["status"]
        if status in ("READY", "ALREADY_EXISTS", "CATEGORY_UPDATE"): counts["valid"] += 1
        if status == "READY": counts["new"] += 1
        elif status == "ALREADY_EXISTS": counts["existing"] += 1
        elif status == "CATEGORY_UPDATE": counts["category_updates"] += 1
        elif status == "DUPLICATE_IN_FILE": counts["duplicates"] += 1; counts["skipped"] += 1
        elif status in ("INVALID_NUMBER", "MISSING_NUMBER"): counts["invalid"] += 1; counts["skipped"] += 1
    return {"summary": counts, "rows": rows[:500]}


@transaction.atomic
def commit(token, user, organization, sheet_name, mappings, default_category_id, confirmed):
    organization = Organization.objects.select_for_update().get(pk=organization.pk)
    value, rows = processed_rows(token, user, organization, sheet_name, mappings, default_category_id, confirmed)
    outcomes = []
    for row in rows:
        if row["status"] in ("DUPLICATE_IN_FILE", "INVALID_NUMBER", "MISSING_NUMBER"): outcomes.append(row); continue
        contact = Contact.objects.select_for_update().filter(organization=organization, phone_number=row["number"]).first()
        if contact is None: contact = next((c for c in Contact.objects.select_for_update().filter(organization=organization) if normalized_phone(c.phone_number) == row["number"]), None)
        created = contact is None; updated = False
        category = ContactCategory.objects.get(pk=row["category_id"], organization=organization)
        if created:
            note_lines = ["Imported from an authorized Excel WhatsApp contact list."]
            if row["website"]: note_lines.append(f'Website: {row["website"]}')
            if row["address"]: note_lines.append(f'Address: {row["address"]}')
            contact = Contact.objects.create(organization=organization, created_by=user, full_name=row["business_name"], company_name=row["business_name"], phone_number=row["number"], email=row["email"], category=category, source=Contact.SOURCE_IMPORT, status=Contact.STATUS_NEW, notes="\n".join(note_lines))
        else:
            fields = []
            for field, incoming in (("full_name", row["business_name"]), ("company_name", row["business_name"]), ("email", row["email"])):
                if not getattr(contact, field) and incoming: setattr(contact, field, incoming); fields.append(field)
            if (not contact.category or contact.category.name == "Unknown") and category.name != "Unknown": contact.category = category; fields.append("category"); updated = True
            if fields: contact.save(update_fields=fields + ["updated_at"])
        key = hashlib.sha256(f'{value["filename"]}\0{sheet_name}\0{row["row"]}\0{row["number"]}'.encode()).hexdigest()
        evidence_status = "CONFIRMED_PUBLIC" if _norm(row["whatsapp_status"]).replace(" ", "_").upper() == "CONFIRMED_PUBLIC" else "USER_CONFIRMED"
        WhatsAppContactImportRecord.objects.get_or_create(organization=organization, source="excel_import", import_key=key, defaults={"contact": contact, "normalized_number": row["number"], "evidence_type": row["evidence_type"] or ("user_declared_excel" if evidence_status == "USER_CONFIRMED" else ""), "evidence_url": row["evidence_url"], "status": evidence_status, "provider_source": row["provider"], "workbook_filename": value["filename"], "worksheet_name": sheet_name, "source_row": row["row"]})
        row["status"] = "CREATED" if created else "CATEGORY_UPDATED" if updated else "ALREADY_EXISTS"; row["contact_id"] = contact.id; outcomes.append(row)
    return {"summary": {"created": sum(r["status"] == "CREATED" for r in outcomes), "existing": sum(r["status"] == "ALREADY_EXISTS" for r in outcomes), "category_updated": sum(r["status"] == "CATEGORY_UPDATED" for r in outcomes), "duplicates": sum(r["status"] == "DUPLICATE_IN_FILE" for r in outcomes), "invalid": sum(r["status"] in ("INVALID_NUMBER", "MISSING_NUMBER") for r in outcomes)}, "rows": outcomes[:500]}
