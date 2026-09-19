import asyncio
import sys
from datetime import datetime, timezone
from io import BytesIO
from unittest.mock import Mock, patch

from django.test import SimpleTestCase
from rest_framework.test import APITestCase

from accounts.models import User
from contacts.models import Contact
from conversations.models import Conversation, Message
from leads.models import Lead, Service
from organizations.models import Organization, OrganizationMember

from lead_collector.models import DiscoveredBusiness, DiscoverySearch, LeadImportRecord
from lead_collector.services.identity import business_identity
from lead_collector.services.providers.discovery_planner import build_query_plan, generate_grid_points
from lead_collector.services.providers.orchestrator import deduplicate_businesses
from lead_collector.services.providers.search_depth import SEARCH_DEPTHS, get_search_depth_profile
from lead_collector.services.providers.playwright_provider import PlaywrightProvider
from lead_collector.services.providers.openstreetmap_provider import OpenStreetMapProvider
from lead_collector.services.enrichment.hotel_website import _extract_contacts
from lead_collector.services.scraping.browser import _start_playwright
from lead_collector.views import export_filename


class DiscoveryAlgorithmTests(SimpleTestCase):
    def test_export_filename_uses_safe_location_slug_and_fallback(self):
        exported_at = datetime(2026, 8, 13, tzinfo=timezone.utc)
        self.assertEqual(export_filename("Manohar International Airport", "xlsx", exported_at), "business-leads-manohar-international-airport-2026-08-13.xlsx")
        self.assertEqual(export_filename("Pune International Airport", "csv", exported_at), "business-leads-pune-international-airport-2026-08-13.csv")
        self.assertEqual(export_filename("Kondapalli, Andhra Pradesh", "xlsx", exported_at), "business-leads-kondapalli-andhra-pradesh-2026-08-13.xlsx")
        self.assertEqual(export_filename('AUX: unsafe? / location*', "csv", exported_at), "business-leads-aux-unsafe-location-2026-08-13.csv")
        self.assertEqual(export_filename("", "xlsx", exported_at), "business-leads-2026-08-13.xlsx")

    @patch("lead_collector.services.scraping.browser.sys.platform", "win32")
    def test_playwright_start_uses_proactor_then_restores_daphne_policy(self):
        initial_policy = asyncio.get_event_loop_policy()
        original_policy = asyncio.WindowsSelectorEventLoopPolicy()
        started = Mock()
        factory = Mock()
        factory.return_value.start.side_effect = lambda: (
            self.assertIsInstance(
                asyncio.get_event_loop_policy(), asyncio.WindowsProactorEventLoopPolicy
            ) or started
        )
        with patch("lead_collector.services.scraping.browser.asyncio.get_event_loop_policy", wraps=asyncio.get_event_loop_policy), patch("lead_collector.services.scraping.browser.asyncio.set_event_loop_policy", wraps=asyncio.set_event_loop_policy):
            asyncio.set_event_loop_policy(original_policy)
            try:
                self.assertIs(_start_playwright(factory), started)
                self.assertIs(asyncio.get_event_loop_policy(), original_policy)
            finally:
                asyncio.set_event_loop_policy(initial_policy)

    def test_all_server_controlled_depths_are_bounded(self):
        self.assertEqual(SEARCH_DEPTHS, ("quick", "standard", "deep", "maximum"))
        expected = {
            "quick": (35, 1, 1, 1, 15, 20, 20, 5),
            "standard": (150, 5, 3, 12, 20, 40, 40, 8),
            "deep": (240, 9, 5, 24, 25, 70, 70, 12),
            "maximum": (360, 9, 7, 48, 30, 100, 100, 16),
        }
        for depth, values in expected.items():
            profile = get_search_depth_profile(depth)
            self.assertEqual((profile.time_budget_seconds, profile.max_grid_points, profile.max_terms, profile.max_queries, profile.results_per_query, profile.max_candidates, profile.max_final_results, profile.scroll_cycles), values)

    def test_quick_query_plan_exhausts_without_waiting_for_time_budget(self):
        profile = get_search_depth_profile("quick")
        points = generate_grid_points(13, 80, 5000, max_points=profile.max_grid_points)
        plan = build_query_plan(["hotels", "resorts"], points, max_terms=profile.max_terms, max_queries=profile.max_queries)
        self.assertEqual(len(points), 1)
        self.assertEqual([(item.term, item.is_center) for item in plan], [("hotels", True)])

    def test_grid_and_center_first_query_plan(self):
        points = generate_grid_points(13, 80, 5000, max_points=5)
        plan = build_query_plan(["hotels", "resorts"], points, max_terms=2, max_queries=8)
        self.assertEqual(len(points), 5)
        self.assertTrue(plan[0].is_center)
        self.assertEqual(plan[0].term, "hotels")

    def test_deduplication_uses_domain_but_not_name_alone(self):
        records = [
            {"name": "Acme", "website": "https://acme.example", "source": "OpenStreetMap", "sources": ["OpenStreetMap"], "_provider": "openstreetmap"},
            {"name": "Acme Hotel", "website": "https://www.acme.example/contact", "source": "Browser Search", "sources": ["Browser Search"], "_provider": "playwright"},
            {"name": "Acme", "address": "Other city", "latitude": 20, "longitude": 20, "source": "Browser Search", "sources": ["Browser Search"], "_provider": "playwright"},
        ]
        self.assertEqual(len(deduplicate_businesses(records)), 2)

    def test_stable_identity_prefers_place_id(self):
        self.assertEqual(business_identity({"name": "A", "place_id": "p1", "source": "Google Places"}), "place:Google Places:p1")

    def test_browser_provider_applies_exact_final_radius(self):
        class FakeScraper:
            last_search_counts = {}
            def search_many(self, **kwargs):
                return [
                    {"name": "Inside", "latitude": 13, "longitude": 80},
                    {"name": "Outside", "latitude": 14, "longitude": 80},
                ]
        results = PlaywrightProvider(FakeScraper()).search_nearby_businesses(13, 80, 5000, "hotels_resorts", "quick")
        self.assertEqual([item["name"] for item in results], ["Inside"])

    @patch("lead_collector.services.providers.openstreetmap_provider.search_nearby_businesses")
    def test_openstreetmap_provider_forwards_category_radius_and_coordinates(self, search):
        search.return_value = [{"name": "Cafe"}]
        result = OpenStreetMapProvider().search_nearby_businesses(13, 80, 5000, "cafes")
        self.assertEqual(result[0]["name"], "Cafe")
        search.assert_called_once_with(13, 80, 5000, "cafes")

    def test_json_ld_list_email_is_optional_and_safe(self):
        html = '<script type="application/ld+json">{"@type":"Hotel","email":["info@example.test"]}</script>'
        result, _ = _extract_contacts(html)
        self.assertIsNone(result["email"])


class CollectorAPITests(APITestCase):
    def setUp(self):
        self.user_a = User.objects.create_user(email="a@example.com", password="pass")
        self.user_b = User.objects.create_user(email="b@example.com", password="pass")
        self.org_a = Organization.objects.create(name="A", slug="a", owner=self.user_a)
        self.org_b = Organization.objects.create(name="B", slug="b", owner=self.user_b)
        OrganizationMember.objects.create(organization=self.org_a, user=self.user_a, role="owner", status="active")
        OrganizationMember.objects.create(organization=self.org_b, user=self.user_b, role="owner", status="active")
        self.service = Service.objects.create(organization=self.org_a, name="Sales", code="SALES")
        self.search_a = DiscoverySearch.objects.create(organization=self.org_a, created_by=self.user_a, location_name="Chennai", latitude=13, longitude=80, radius_m=5000, category="hotels_resorts", providers=["openstreetmap"], status="complete")
        self.business_a = DiscoveredBusiness.objects.create(organization=self.org_a, search=self.search_a, identity_key="place:osm:1", name="Public Hotel", phone_normalized="+919876543210", payload={"name": "Public Hotel", "phone": "+91 98765 43210", "address": "Chennai", "website": "https://hotel.example"})

    def test_authentication_is_required(self):
        self.assertEqual(self.client.get("/api/lead-collector/categories/").status_code, 401)

    def test_owner_can_save_provider_key_without_api_disclosure(self):
        self.client.force_authenticate(self.user_a)
        response = self.client.patch("/api/lead-collector/providers/", {
            "business_providers": ["openstreetmap", "geoapify"],
            "people_providers": ["official_website"],
            "geoapify_api_key": "organization-a-secret",
        }, format="json")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.data["providers"]["geoapify"]["configured"])
        self.assertNotContains(response, "organization-a-secret")
        self.assertNotIn("geoapify_api_key", response.data)

        self.client.force_authenticate(self.user_b)
        other = self.client.get("/api/lead-collector/providers/")
        self.assertFalse(other.data["providers"]["geoapify"]["configured"])

    def test_agent_cannot_change_provider_settings(self):
        agent = User.objects.create_user(email="agent@example.com", password="pass")
        OrganizationMember.objects.create(organization=self.org_a, user=agent, role="agent", status="active")
        self.client.force_authenticate(agent)
        response = self.client.patch("/api/lead-collector/providers/", {
            "business_providers": ["openstreetmap"], "people_providers": []
        }, format="json")
        self.assertEqual(response.status_code, 403)

    def test_keyed_provider_requires_credential_before_enablement(self):
        self.client.force_authenticate(self.user_a)
        response = self.client.patch("/api/lead-collector/providers/", {
            "business_providers": ["openstreetmap", "google"],
            "people_providers": ["official_website"],
        }, format="json")
        self.assertEqual(response.status_code, 400)
        self.assertIn("Google Places", response.data["detail"])

    @patch("lead_collector.views.search_locations")
    def test_location_search_is_authenticated_and_normalized(self, search):
        search.return_value = [{"display_name": "Chennai, India", "latitude": 13, "longitude": 80}]
        self.client.force_authenticate(self.user_a)
        response = self.client.get("/api/lead-collector/locations/?q=chen")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["results"][0]["latitude"], 13)

    @patch("lead_collector.views.search_locations")
    def test_short_location_query_does_not_call_provider(self, search):
        self.client.force_authenticate(self.user_a)
        response = self.client.get("/api/lead-collector/locations/?q=vi")
        self.assertEqual(response.status_code, 400)
        search.assert_not_called()

    def test_cross_organization_search_and_business_access_is_hidden(self):
        self.client.force_authenticate(self.user_b)
        self.assertEqual(self.client.get(f"/api/lead-collector/searches/{self.search_a.id}/").status_code, 404)
        response = self.client.post("/api/lead-collector/imports/preview/", {"business_ids": [self.business_a.id], "search_id": self.search_a.id}, format="json")
        self.assertEqual(response.status_code, 404)

    @patch("lead_collector.views.search_businesses_multi_provider")
    def test_search_persists_only_in_current_organization(self, search):
        search.return_value = {"hotels": [{"name": "Cafe", "phone": "+919111111111", "address": "Chennai", "source": "OpenStreetMap", "osm_id": "2"}], "provider_results": {"openstreetmap": {"status": "success", "count": 1}}, "raw_result_count": 1, "deduplicated_count": 1}
        self.client.force_authenticate(self.user_a)
        response = self.client.post("/api/lead-collector/searches/", {"location_name": "Chennai", "latitude": 13, "longitude": 80, "radius_m": 5000, "category": "cafes", "providers": ["openstreetmap"], "search_depth": "maximum"}, format="json")
        self.assertEqual(response.status_code, 201)
        self.assertEqual(response.data["search_depth"], "standard")
        self.assertEqual(DiscoveredBusiness.objects.get(name="Cafe").organization, self.org_a)

    def test_import_preview_creation_and_duplicate_prevention_without_sending(self):
        self.client.force_authenticate(self.user_a)
        preview = self.client.post("/api/lead-collector/imports/preview/", {"business_ids": [self.business_a.id], "search_id": self.search_a.id}, format="json")
        self.assertTrue(preview.data["results"][0]["can_import"])
        payload = {"business_ids": [self.business_a.id], "service_id": self.service.id, "search_id": self.search_a.id}
        first = self.client.post("/api/lead-collector/imports/leads/", payload, format="json")
        second = self.client.post("/api/lead-collector/imports/leads/", payload, format="json")
        self.assertEqual(first.data["results"][0]["status"], "created")
        self.assertEqual(second.data["results"][0]["status"], "duplicate")
        lead = Lead.objects.get()
        self.assertEqual(lead.source, "lead_collector")
        self.assertEqual(lead.organization, self.org_a)
        self.assertEqual(lead.phone, "+919876543210")
        self.assertEqual(Contact.objects.count(), 0)
        self.assertEqual(Conversation.objects.count(), 0)
        self.assertEqual(Message.objects.count(), 0)

    def test_import_duplicate_evidence_is_org_scoped_and_preserves_distinct_branches(self):
        same_phone = DiscoveredBusiness.objects.create(
            organization=self.org_a, search=self.search_a, identity_key="place:osm:same-phone",
            name="Other Name", phone_normalized="+919876543210",
            payload={"name": "Other Name", "phone": "+91 98765 43210", "place_id": "different-place", "source": "Google Places"},
        )
        same_place = DiscoveredBusiness.objects.create(
            organization=self.org_a, search=self.search_a, identity_key="place:google:shared",
            name="Provider Match", phone_normalized="+919000000001",
            payload={"name": "Provider Match", "phone": "+91 90000 00001", "place_id": "shared-place", "source": "Google Places"},
        )
        same_place_again = DiscoveredBusiness.objects.create(
            organization=self.org_a, search=self.search_a, identity_key="place:google:shared-second-search-row",
            name="Provider Match Renamed", phone_normalized="+919000000002",
            payload={"name": "Provider Match Renamed", "phone": "+91 90000 00002", "place_id": "shared-place", "source": "Google Places"},
        )
        branch_one = DiscoveredBusiness.objects.create(
            organization=self.org_a, search=self.search_a, identity_key="place:google:branch-one",
            name="Acme", phone_normalized="+919100000001",
            payload={"name": "Acme", "phone": "+91 91000 00001", "place_id": "branch-one", "source": "Google Places", "address": "North"},
        )
        branch_two = DiscoveredBusiness.objects.create(
            organization=self.org_a, search=self.search_a, identity_key="place:google:branch-two",
            name="Acme", phone_normalized="+919100000002",
            payload={"name": "Acme", "phone": "+91 91000 00002", "place_id": "branch-two", "source": "Google Places", "address": "South"},
        )
        self.client.force_authenticate(self.user_a)
        endpoint = "/api/lead-collector/imports/leads/"
        first = self.client.post(endpoint, {"business_ids": [self.business_a.id], "service_id": self.service.id, "search_id": self.search_a.id}, format="json")
        phone_duplicate = self.client.post(endpoint, {"business_ids": [same_phone.id], "service_id": self.service.id, "search_id": self.search_a.id}, format="json")
        place_first = self.client.post(endpoint, {"business_ids": [same_place.id], "service_id": self.service.id, "search_id": self.search_a.id}, format="json")
        place_duplicate = self.client.post(endpoint, {"business_ids": [same_place_again.id], "service_id": self.service.id, "search_id": self.search_a.id}, format="json")
        branches = self.client.post(endpoint, {"business_ids": [branch_one.id, branch_two.id], "service_id": self.service.id, "search_id": self.search_a.id}, format="json")
        self.assertEqual(first.data["results"][0]["status"], "created")
        self.assertEqual(phone_duplicate.data["results"][0]["reason"], "business_phone")
        self.assertEqual(place_first.data["results"][0]["status"], "created")
        self.assertEqual(place_duplicate.data["results"][0]["reason"], "provider_place_id")
        self.assertEqual([row["status"] for row in branches.data["results"]], ["created", "created"])

        service_b = Service.objects.create(organization=self.org_b, name="Sales", code="SALES")
        search_b = DiscoverySearch.objects.create(organization=self.org_b, created_by=self.user_b, location_name="Chennai", latitude=13, longitude=80, radius_m=5000, category="hotels_resorts", providers=["openstreetmap"], status="complete")
        business_b = DiscoveredBusiness.objects.create(organization=self.org_b, search=search_b, identity_key="place:osm:1", name="Public Hotel", phone_normalized="+919876543210", payload={"name": "Public Hotel", "phone": "+91 98765 43210", "place_id": "shared-place", "source": "Google Places"})
        self.client.force_authenticate(self.user_b)
        other_org = self.client.post(endpoint, {"business_ids": [business_b.id], "service_id": service_b.id, "search_id": search_b.id}, format="json")
        self.assertEqual(other_org.data["results"][0]["status"], "created")
        self.assertEqual(Lead.objects.filter(phone="+919876543210").count(), 2)

    def test_import_acquires_org_lock_before_duplicate_check_and_retry_is_idempotent(self):
        from lead_collector.services.crm_import import duplicate_for as real_duplicate_for
        events = []
        self.client.force_authenticate(self.user_a)
        payload = {"business_ids": [self.business_a.id], "service_id": self.service.id, "search_id": self.search_a.id}
        with patch("lead_collector.services.crm_import.Organization.objects.select_for_update") as select_for_update, patch(
            "lead_collector.services.crm_import.duplicate_for",
            side_effect=lambda *args: events.append("duplicate-check") or real_duplicate_for(*args),
        ):
            select_for_update.side_effect = lambda: events.append("organization-lock") or Mock(
                get=Mock(return_value=self.org_a)
            )
            first = self.client.post("/api/lead-collector/imports/leads/", payload, format="json")
        second = self.client.post("/api/lead-collector/imports/leads/", payload, format="json")
        self.assertEqual(events[:2], ["organization-lock", "duplicate-check"])
        self.assertEqual(first.data["results"][0]["status"], "created")
        self.assertEqual(second.data["results"][0]["status"], "duplicate")
        self.assertEqual(Lead.objects.filter(organization=self.org_a).count(), 1)
        self.assertEqual(LeadImportRecord.objects.filter(organization=self.org_a, business=self.business_a).count(), 1)

    def test_invalid_service_and_missing_phone_are_safe(self):
        no_phone = DiscoveredBusiness.objects.create(organization=self.org_a, search=self.search_a, identity_key="name:none", name="No Phone", payload={"name": "No Phone"})
        self.client.force_authenticate(self.user_a)
        invalid_service = self.client.post("/api/lead-collector/imports/leads/", {"business_ids": [self.business_a.id], "service_id": 99999, "search_id": self.search_a.id}, format="json")
        self.assertEqual(invalid_service.status_code, 400)
        result = self.client.post("/api/lead-collector/imports/leads/", {"business_ids": [no_phone.id], "service_id": self.service.id, "search_id": self.search_a.id}, format="json")
        self.assertEqual(result.data["results"][0]["reason"], "missing_phone")

    def test_csv_and_excel_exports_are_scoped_and_formula_safe(self):
        self.business_a.payload.update({"id": "provider-place-id", "name": "=HYPERLINK(\"bad\")", "email": None, "brand": None, "website": "https://hotel.example", "rating": 4.2, "review_count": 0, "enrichment_status": "NOT_FOUND", "business_emails": [], "business_phones": [], "decision_makers": [], "social_profiles": {}, "whatsapp_contacts": []})
        self.business_a.save(update_fields=["payload"])
        self.client.force_authenticate(self.user_a)
        detail = self.client.get(f"/api/lead-collector/searches/{self.search_a.id}/")
        self.assertEqual(detail.data["businesses"][0]["id"], self.business_a.id)
        payload = {"business_ids": [self.business_a.id]}
        csv_response = self.client.post("/api/lead-collector/exports/csv/", payload, format="json")
        excel_response = self.client.post("/api/lead-collector/exports/excel/", payload, format="json")
        self.assertEqual(csv_response.status_code, 200)
        self.assertIn("'=HYPERLINK", csv_response.content.decode())
        self.assertEqual(excel_response.status_code, 200)
        self.assertEqual(excel_response["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertRegex(excel_response["Content-Disposition"], r'business-leads-chennai-\d{4}-\d{2}-\d{2}\.xlsx')

    def test_excel_export_supports_multiple_real_records_and_keeps_foreign_records_hidden(self):
        self.business_a.payload.update({"enrichment_status": "NOT_FOUND", "business_emails": [], "business_phones": [], "decision_makers": [], "social_profiles": {}})
        self.business_a.save(update_fields=["payload"])
        second = DiscoveredBusiness.objects.create(organization=self.org_a, search=self.search_a, identity_key="place:provider:2", name="Unicode Café", payload={"id": "provider-two", "name": "Unicode Café", "phone": "+91 90000 00000", "email": None, "website": None, "decision_makers": [], "social_profiles": {}, "whatsapp_contacts": []})
        foreign_search = DiscoverySearch.objects.create(organization=self.org_b, created_by=self.user_b, location_name="Other", latitude=1, longitude=2, radius_m=1000, category="cafes", providers=["openstreetmap"], status="complete")
        foreign = DiscoveredBusiness.objects.create(organization=self.org_b, search=foreign_search, identity_key="foreign", name="Foreign", payload={"name": "Foreign"})
        self.client.force_authenticate(self.user_a)
        response = self.client.post("/api/lead-collector/exports/excel/", {"business_ids": [self.business_a.id, second.id, foreign.id]}, format="json")
        self.assertEqual(response.status_code, 200)
        from openpyxl import load_workbook
        from io import BytesIO
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        self.assertEqual(set(workbook.sheetnames), {"Search Summary", "Businesses", "Business Contacts", "Decision Makers", "Social Profiles", "WhatsApp Contacts"})
        rows = list(workbook["Businesses"].iter_rows(values_only=True))
        name_column = rows[0].index("Business Name")
        names = [row[name_column] for row in rows[1:]]
        self.assertIn("Unicode Café", names)
        self.assertNotIn("Foreign", names)

    @patch("lead_collector.services.export.enrichment.enrich_hotels_bulk")
    def test_export_prepares_only_eligible_records_persists_enrichment_and_isolates_failures_and_orgs(self, enrich_bulk):
        already = DiscoveredBusiness.objects.create(
            organization=self.org_a, search=self.search_a, identity_key="already", name="Already",
            payload={"name": "Already", "website": "https://already.example", "enrichment_status": "FOUND", "business_emails": [], "business_phones": [], "social_profiles": {}, "decision_makers": []},
        )
        no_website = DiscoveredBusiness.objects.create(
            organization=self.org_a, search=self.search_a, identity_key="no-site", name="No Site",
            payload={"name": "No Site"},
        )
        foreign_search = DiscoverySearch.objects.create(organization=self.org_b, created_by=self.user_b, location_name="Foreign", latitude=1, longitude=2, radius_m=1000, category="cafes", providers=["openstreetmap"], status="complete")
        foreign = DiscoveredBusiness.objects.create(organization=self.org_b, search=foreign_search, identity_key="foreign-export", name="Foreign", payload={"name": "Foreign", "website": "https://foreign.example"})
        enrich_bulk.return_value = {"results": [{
            "status": "FOUND", "hotel": {**self.business_a.payload, "email": "hello@hotel.example"},
            "sources": {"email": "https://hotel.example/contact"}, "website_confidence": "HIGH",
            "business_emails": [{"email": "hello@hotel.example", "type": "general", "source_url": "https://hotel.example/contact"}],
            "business_phones": [{"phone": "+91 98765 43210", "normalized": "+919876543210", "type": "general", "source_url": "https://hotel.example/contact"}],
            "whatsapp_contacts": [{"number": "+91 98765 43210", "normalized": "+919876543210", "status": "CONFIRMED_PUBLIC", "evidence_type": "wa.me link", "source_url": "https://hotel.example/contact"}],
            "social_profiles": {"linkedin": "https://linkedin.com/company/hotel", "facebook": "https://facebook.com/hotel", "instagram": "https://instagram.com/hotel"},
            "social_profile_sources": {"linkedin": "https://hotel.example", "facebook": "https://hotel.example", "instagram": "https://hotel.example"},
            "discovered_pages": {"contact": "https://hotel.example/contact"},
            "decision_makers": [{"name": "Jane Public", "title": "General Manager", "role_group": "general_manager", "email": "jane@hotel.example", "source": "Official Website", "source_url": "https://hotel.example/team"}],
        }]}
        self.client.force_authenticate(self.user_a)
        response = self.client.post("/api/lead-collector/exports/excel/", {
            "business_ids": [self.business_a.id, already.id, no_website.id, foreign.id],
            "context": {"id": self.search_a.id, "location_name": "stale client value"},
        }, format="json")
        self.assertEqual(response.status_code, 200)
        enrich_bulk.assert_called_once()
        self.assertEqual([item["name"] for item in enrich_bulk.call_args.args[0]], ["Public Hotel"])
        self.business_a.refresh_from_db(); foreign.refresh_from_db()
        self.assertEqual(self.business_a.payload["email"], "hello@hotel.example")
        self.assertNotIn("enrichment_status", foreign.payload)
        from openpyxl import load_workbook
        workbook = load_workbook(BytesIO(response.content), read_only=True)
        self.assertEqual(list(workbook["Business Contacts"].values)[1][2], "hello@hotel.example")
        self.assertEqual(list(workbook["Decision Makers"].values)[1][1], "Jane Public")
        self.assertEqual(list(workbook["Social Profiles"].values)[1][1], "https://linkedin.com/company/hotel")
        self.assertEqual(list(workbook["WhatsApp Contacts"].values)[1][4], "CONFIRMED_PUBLIC")
        self.assertIn("business-leads-chennai-", response["Content-Disposition"])
        csv_response = self.client.post("/api/lead-collector/exports/csv/", {
            "business_ids": [self.business_a.id], "context": {"id": self.search_a.id},
        }, format="json")
        self.assertEqual(csv_response.status_code, 200)
        csv_text = csv_response.content.decode()
        self.assertIn("hello@hotel.example", csv_text)
        self.assertIn("CONFIRMED_PUBLIC", csv_text)
        self.assertIn("https://linkedin.com/company/hotel", csv_text)
        self.assertIn("business-leads-chennai-", csv_response["Content-Disposition"])
        enrich_bulk.assert_called_once()

    @patch("lead_collector.views.search_decision_makers_multi_provider")
    def test_decision_maker_search_is_scoped_and_persisted(self, search_people):
        search_people.return_value = {
            "contacts": [{"name": "Jane Public", "title": "General Manager", "source": "Official Website"}],
            "status": "FOUND", "provider_results": {"official_website": {"status": "success"}},
        }
        self.client.force_authenticate(self.user_a)
        response = self.client.post("/api/lead-collector/businesses/decision-makers/", {
            "business_ids": [self.business_a.id], "providers": ["official_website"], "search_id": self.search_a.id
        }, format="json")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.data["businesses"][0]["decision_makers"][0]["name"], "Jane Public")
        self.business_a.refresh_from_db()
        self.assertEqual(self.business_a.payload["decision_maker_status"], "FOUND")

        self.client.force_authenticate(self.user_b)
        hidden = self.client.post("/api/lead-collector/businesses/decision-makers/", {
            "business_ids": [self.business_a.id], "providers": ["official_website"], "search_id": self.search_a.id
        }, format="json")
        self.assertEqual(hidden.status_code, 404)
